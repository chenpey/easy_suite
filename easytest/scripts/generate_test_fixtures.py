from __future__ import annotations

import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw

from easytest.cases.compiler import compile_workbook

ROOT = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "project"
CASES_DIR = ROOT / "cases"
ASSET_DIR = ROOT / "assets"
FIXED_TIME = datetime(2026, 1, 1, tzinfo=UTC)

CASE_HEADERS = [
    "case_id",
    "case_name",
    "case_type",
    "enabled",
    "tags",
    "variables",
    "mock_profile",
    "snapshot_profile",
]
STEP_HEADERS = [
    "case_id",
    "step_id",
    "order",
    "executor",
    "operation",
    "request",
    "save_as",
    "mock",
    "snapshot",
    "expect",
]


def _json(value: Any) -> str:
    if value in (None, ""):
        return ""
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _write_bytes(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        with path.open("r+b") as stream:
            stream.seek(0)
            stream.write(content)
            stream.truncate()
    else:
        path.write_bytes(content)


def _workbook(path: Path, cases: list[list[Any]], steps: list[list[Any]]) -> None:
    workbook = Workbook()
    workbook.properties.creator = "easytest"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    cases_sheet = workbook.active
    cases_sheet.title = "cases"
    steps_sheet = workbook.create_sheet("steps")

    for sheet, headers, rows in (
        (cases_sheet, CASE_HEADERS, cases),
        (steps_sheet, STEP_HEADERS, steps),
    ):
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for column in sheet.columns:
            width = min(
                60,
                max(12, max(len(str(cell.value or "")) for cell in column) + 2),
            )
            sheet.column_dimensions[column[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    stream = BytesIO()
    workbook.save(stream)
    _write_bytes(path, stream.getvalue())
    compile_workbook(path)


def _dashboard_image() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (640, 360), "#f5f5f7")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle(
        (24, 24, 616, 336), radius=18, fill="#ffffff", outline="#d2d2d7"
    )
    draw.text((52, 52), "Account Dashboard", fill="#1d1d1f")
    draw.text((52, 92), "Mock User", fill="#6e6e73")
    draw.rounded_rectangle((52, 132, 300, 250), radius=10, fill="#0071e3")
    draw.text((72, 154), "Available limit", fill="#ffffff")
    draw.text((72, 192), "1,000.00", fill="#ffffff")
    draw.rounded_rectangle((324, 132, 568, 250), radius=10, fill="#e8f5e9")
    draw.text((344, 154), "Status", fill="#3a3a3c")
    draw.text((344, 192), "ACTIVE", fill="#188038")
    stream = BytesIO()
    image.save(stream, format="PNG", optimize=True)
    _write_bytes(ASSET_DIR / "dashboard.png", stream.getvalue())


def main() -> None:
    _dashboard_image()
    _workbook(
        CASES_DIR / "scenario" / "account_flow.xlsx",
        [
            [
                "scenario.account_flow",
                "Account flow across adapters",
                "scenario",
                True,
                "smoke,offline",
                _json({"requested_amount": "500.00"}),
                "",
                "default",
            ]
        ],
        [
            [
                "scenario.account_flow",
                "seed_user",
                1,
                "scenario",
                "scenario.set_state",
                _json({"path": "user.id", "value": "user-001"}),
                "",
                "",
                "",
                _json({"path": "$.value", "equals": "user-001"}),
            ],
            [
                "scenario.account_flow",
                "profile",
                2,
                "http",
                "http.get_profile",
                _json({"params": {"id": "${state.user.id}"}}),
                "profile",
                "",
                _json({"name": "profile_response", "rule": "response_default"}),
                _json({"path": "$.status_code", "equals": 200}),
            ],
            [
                "scenario.account_flow",
                "limit",
                3,
                "rpc",
                "rpc.calculate_limit",
                _json(
                    {
                        "amount": "${variables.requested_amount}",
                        "user_id": "${steps.profile.body.id}",
                    }
                ),
                "",
                "",
                _json({"name": "limit_response", "rule": "response_default"}),
                _json({"path": "$.approved", "equals": True}),
            ],
            [
                "scenario.account_flow",
                "accounts",
                4,
                "database",
                "database.list_accounts",
                _json({"parameters": {"user_id": "${state.user.id}"}}),
                "",
                "",
                _json({"name": "accounts", "rule": "database_rows"}),
                _json({"path": "$.rowcount", "equals": 2}),
            ],
            [
                "scenario.account_flow",
                "dashboard",
                5,
                "ui",
                "ui.open_dashboard",
                _json({"user_id": "${state.user.id}"}),
                "",
                "",
                _json({"name": "dashboard", "rule": "screenshot_default"}),
                _json({"path": "$.visible", "equals": True}),
            ],
        ],
    )
    _workbook(
        CASES_DIR / "http" / "profile_api.xlsx",
        [
            [
                "http.profile_api",
                "Profile HTTP API",
                "http",
                True,
                "api,offline",
                _json({"user_id": "user-001"}),
                "",
                "default",
            ]
        ],
        [
            [
                "http.profile_api",
                "get_profile",
                1,
                "http",
                "http.get_profile",
                _json({"params": {"id": "${variables.user_id}"}}),
                "",
                "",
                _json({"name": "response", "rule": "response_default"}),
                _json(
                    {
                        "checks": [
                            {"path": "$.status_code", "equals": 200},
                            {"path": "$.body.status", "equals": "ACTIVE"},
                        ]
                    }
                ),
            ]
        ],
    )
    _workbook(
        CASES_DIR / "rpc" / "limit_service.xlsx",
        [
            [
                "rpc.limit_service",
                "Limit RPC service",
                "rpc",
                True,
                "rpc,offline",
                _json({"amount": "500.00", "user_id": "user-001"}),
                "",
                "default",
            ]
        ],
        [
            [
                "rpc.limit_service",
                "calculate_limit",
                1,
                "rpc",
                "rpc.calculate_limit",
                _json(
                    {
                        "amount": "${variables.amount}",
                        "user_id": "${variables.user_id}",
                    }
                ),
                "",
                "",
                _json({"name": "response", "rule": "response_default"}),
                _json({"path": "$.approved", "equals": True}),
            ]
        ],
    )


if __name__ == "__main__":
    main()
