from __future__ import annotations

import hashlib
import json
import os
import re
import stat
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from easytest._version import __version__
from easytest.cases.schema import SCHEMA_VERSION, parse_document
from easytest.models import ContractError

CASE_COLUMNS = {"case_id", "case_name", "case_type"}
STEP_COLUMNS = {"case_id", "step_id", "order", "executor", "operation"}
DATA_COLUMNS = {"data_set", "data_id"}
CASE_OPTIONAL_COLUMNS = {
    "enabled",
    "tags",
    "variables",
    "mock_profile",
    "snapshot_profile",
    "data_set",
}
STEP_OPTIONAL_COLUMNS = {"request", "save_as", "mock", "snapshot", "expect"}
JSON_COLUMNS = {"variables", "request", "mock", "snapshot", "expect"}
_DATA_FIELD_PATTERN = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def _location(error: ContractError, path: Path, sheet: str, row: int,
              column: str, value: Any = None, expected: str = "") -> ContractError:
    error.easytest_location = {
        "source": str(path), "sheet": sheet, "source_row": row, "column": column,
        "actual_type": type(value).__name__, "expected_type": expected,
    }
    error.add_note(
        f"Source={path}, sheet={sheet}, row={row}, column={column}, "
        f"actual_type={type(value).__name__}, expected_type={expected}"
    )
    return error


def _atomic_write(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            dir=path.parent,
            prefix=f".{path.name}.",
            delete=False,
        ) as stream:
            temporary = Path(stream.name)
            stream.write(content)
        mode = stat.S_IMODE(path.stat().st_mode) if path.exists() else 0o644
        os.chmod(temporary, mode)
        os.replace(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _parse_bool(value: Any, *, default: bool = True) -> bool:
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "y", "on"}:
        return True
    if normalized in {"0", "false", "no", "n", "off"}:
        return False
    raise ContractError(f"invalid boolean value: {value!r}")


def _parse_json_cell(value: Any, *, field: str) -> Any:
    if value in (None, ""):
        return None
    if isinstance(value, (dict, list, bool, int, float)):
        return value
    text = str(value).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise ContractError(f"{field} must contain valid JSON: {exc.msg}") from exc


def _cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def _sheet_rows(workbook_path: Path, worksheet) -> list[dict[str, Any]]:
    cells = iter(worksheet.iter_rows())
    first = next(cells, None)
    if first is None:
        raise ContractError(f"{workbook_path.name}:{worksheet.title} is empty")

    headers = [
        str(cell.value).strip() if cell.value is not None else "" for cell in first
    ]
    if any(not header for header in headers):
        raise ContractError(
            f"{workbook_path.name}:{worksheet.title} has a blank header"
        )
    if len(headers) != len(set(headers)):
        raise ContractError(
            f"{workbook_path.name}:{worksheet.title} has duplicate headers"
        )
    allowed = (
        CASE_COLUMNS | CASE_OPTIONAL_COLUMNS if worksheet.title == "cases"
        else STEP_COLUMNS | STEP_OPTIONAL_COLUMNS
    )
    unknown = set(headers) - allowed
    if unknown:
        raise ContractError(
            f"{workbook_path.name}:{worksheet.title} row=1 has unknown columns: {sorted(unknown)}"
        )

    rows: list[dict[str, Any]] = []
    for row_index, row in enumerate(cells, start=2):
        if all(cell.value in (None, "") for cell in row):
            continue
        record: dict[str, Any] = {}
        for column_index, (header, cell) in enumerate(
            zip(headers, row, strict=False),
            start=1,
        ):
            coordinate = getattr(
                cell,
                "coordinate",
                f"{get_column_letter(column_index)}{row_index}",
            )
            if cell.data_type == "f":
                raise ContractError(
                    f"{workbook_path.name}:{worksheet.title}!{coordinate} "
                    "contains a formula; use a literal value"
                )
            if header in {
                "case_id", "case_name", "case_type", "step_id", "executor",
                "operation", "save_as", "mock_profile", "snapshot_profile", "tags",
                "data_set",
            } and cell.value is not None and not isinstance(cell.value, str):
                raise _location(
                    ContractError(f"{header} must be a string; automatic Excel type conversion is not supported",
                                  code="INVALID_TYPE", field=header),
                    workbook_path, worksheet.title, row_index, header, cell.value, "str",
                )
            value = _cell_value(cell.value)
            if header in JSON_COLUMNS:
                try:
                    value = _parse_json_cell(
                        value,
                        field=f"{workbook_path.name}:{worksheet.title}!{coordinate}",
                    )
                    if header in {"variables", "request"} and value is not None and not isinstance(value, dict):
                        raise ContractError(f"{header} must be a JSON object", code="INVALID_TYPE", field=header)
                except ContractError as error:
                    raise _location(error, workbook_path, worksheet.title, row_index,
                                    header, value, "JSON object" if header in {"variables", "request"} else "JSON")
            if header == "case_type" and (not isinstance(value, str) or value.lower() not in {"scenario", "http", "rpc"}):
                raise _location(
                    ContractError("case type must be one of scenario/http/rpc", field="case_type"),
                    workbook_path, worksheet.title, row_index, header, cell.value, "scenario/http/rpc",
                )
            if header == "enabled":
                try:
                    value = _parse_bool(value)
                except ContractError as error:
                    raise _location(error, workbook_path, worksheet.title, row_index,
                                    header, cell.value, "boolean")
            record[header] = value
        record["_row"] = row_index
        rows.append(record)
    return rows


def _data_sheet_rows(
    workbook_path: Path,
    worksheet,
) -> tuple[list[str], list[dict[str, Any]]]:
    cells = iter(worksheet.iter_rows())
    first = next(cells, None)
    if first is None:
        raise ContractError(f"{workbook_path.name}:data is empty")
    headers = [
        str(cell.value).strip() if cell.value is not None else "" for cell in first
    ]
    if any(not header for header in headers):
        raise ContractError(f"{workbook_path.name}:data has a blank header")
    if len(headers) != len(set(headers)):
        raise ContractError(f"{workbook_path.name}:data has duplicate headers")
    missing = DATA_COLUMNS - set(headers)
    if missing:
        raise ContractError(
            f"{workbook_path.name}:data is missing columns: {sorted(missing)}"
        )
    value_columns = [
        header for header in headers if header not in DATA_COLUMNS | {"enabled"}
    ]
    if not value_columns:
        raise ContractError(
            f"{workbook_path.name}:data must contain at least one business column"
        )
    invalid = [
        header for header in value_columns if not _DATA_FIELD_PATTERN.fullmatch(header)
    ]
    if invalid:
        raise _location(
            ContractError(
                f"{workbook_path.name}:data has invalid business columns: "
                f"{invalid}; use Python-style identifiers",
                code="INVALID_IDENTIFIER",
                field=invalid[0],
            ),
            workbook_path,
            "data",
            1,
            invalid[0],
            invalid[0],
            "Python-style identifier",
        )

    rows: list[dict[str, Any]] = []
    for row_index, row in enumerate(cells, start=2):
        if all(cell.value in (None, "") for cell in row):
            continue
        record: dict[str, Any] = {}
        for column_index, (header, cell) in enumerate(
            zip(headers, row, strict=False),
            start=1,
        ):
            coordinate = getattr(
                cell,
                "coordinate",
                f"{get_column_letter(column_index)}{row_index}",
            )
            if cell.data_type == "f":
                raise ContractError(
                    f"{workbook_path.name}:data!{coordinate} "
                    "contains a formula; use a literal value"
                )
            if (
                header in DATA_COLUMNS
                and cell.value is not None
                and not isinstance(cell.value, str)
            ):
                raise _location(
                    ContractError(
                        f"{header} must be a string; automatic Excel type "
                        "conversion is not supported",
                        code="INVALID_TYPE",
                        field=header,
                    ),
                    workbook_path,
                    "data",
                    row_index,
                    header,
                    cell.value,
                    "str",
                )
            value = _cell_value(cell.value)
            if header == "enabled":
                try:
                    value = _parse_bool(value)
                except ContractError as error:
                    raise _location(
                        error,
                        workbook_path,
                        "data",
                        row_index,
                        header,
                        cell.value,
                        "boolean",
                    )
            record[header] = value
        record["_row"] = row_index
        rows.append(record)
    if not rows:
        raise ContractError(f"{workbook_path.name}:data must contain at least one data row")
    return value_columns, rows


def _require_columns(
    path: Path, sheet: str, rows: list[dict[str, Any]], required: set[str]
) -> None:
    if not rows:
        raise ContractError(f"{path.name}:{sheet} must contain at least one data row")
    missing = required - set(rows[0])
    if missing:
        raise ContractError(
            f"{path.name}:{sheet} is missing columns: {sorted(missing)}"
        )


def workbook_document(path: str | Path) -> dict[str, Any]:
    workbook_path = Path(path).resolve()
    if workbook_path.suffix.lower() != ".xlsx":
        raise ContractError(f"only .xlsx case sources are supported: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        missing_sheets = {"cases", "steps"} - set(workbook.sheetnames)
        if missing_sheets:
            raise ContractError(
                f"{workbook_path.name} is missing sheets: {sorted(missing_sheets)}"
            )
        case_rows = _sheet_rows(workbook_path, workbook["cases"])
        step_rows = _sheet_rows(workbook_path, workbook["steps"])
        has_data_sheet = "data" in workbook.sheetnames
        data_columns: list[str] = []
        data_rows: list[dict[str, Any]] = []
        if has_data_sheet:
            data_columns, data_rows = _data_sheet_rows(
                workbook_path,
                workbook["data"],
            )
    finally:
        workbook.close()

    _require_columns(workbook_path, "cases", case_rows, CASE_COLUMNS)
    _require_columns(workbook_path, "steps", step_rows, STEP_COLUMNS)

    steps_by_case: dict[str, list[dict[str, Any]]] = {}
    for row in step_rows:
        case_id = str(row.get("case_id") or "").strip()
        step = {
            "id": row.get("step_id"),
            "order": row.get("order"),
            "executor": row.get("executor"),
            "operation": row.get("operation"),
            "request": row.get("request"),
            "save_as": row.get("save_as") or None,
            "mock": row.get("mock"),
            "snapshot": row.get("snapshot"),
            "expect": row.get("expect"),
            "source_row": row["_row"],
        }
        steps_by_case.setdefault(case_id, []).append(step)

    data_sets: dict[str, list[dict[str, Any]]] = {}
    for row in data_rows:
        data_set = str(row.get("data_set") or "").strip()
        data_sets.setdefault(data_set, []).append(
            {
                "id": row.get("data_id"),
                "enabled": _parse_bool(row.get("enabled")),
                "values": {
                    column: row.get(column)
                    for column in data_columns
                },
                "source_row": row["_row"],
            }
        )

    cases: list[dict[str, Any]] = []
    seen_case_ids: set[str] = set()
    for row in case_rows:
        case_id = str(row.get("case_id") or "").strip()
        if case_id in seen_case_ids:
            raise _location(
                ContractError(f"duplicate case_id {case_id!r} in cases sheet", field="case_id"),
                workbook_path, "cases", row["_row"], "case_id", row.get("case_id"), "unique string",
            )
        seen_case_ids.add(case_id)
        case = {
            "id": case_id,
            "name": row.get("case_name") or case_id,
            "type": str(row.get("case_type") or "").lower(),
            "enabled": _parse_bool(row.get("enabled")),
            "tags": row.get("tags") or "",
            "variables": row.get("variables"),
            "mock_profile": row.get("mock_profile") or None,
            "snapshot_profile": row.get("snapshot_profile") or "default",
            "steps": steps_by_case.pop(case_id, []),
        }
        if row.get("data_set"):
            case["data_set"] = row["data_set"]
        cases.append(case)
    if steps_by_case:
        raise ContractError(
            f"steps reference unknown case ids: {sorted(steps_by_case)}"
        )

    document = {
        "schema_version": SCHEMA_VERSION,
        "source_mode": "xlsx",
        "source": workbook_path.name,
        "source_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "compiler_version": __version__,
        "cases": sorted(cases, key=lambda item: item["id"]),
    }
    if has_data_sheet:
        document["data_sets"] = {
            name: data_sets[name]
            for name in sorted(data_sets)
        }
    try:
        parse_document(document, source=str(workbook_path))
    except ContractError as error:
        location = getattr(error, "easytest_location", {})
        if location.get("source_row") is None:
            case_id = location.get("case_id")
            row = next((item for item in case_rows if item.get("case_id") == case_id), None)
            if row is not None:
                column = {
                    "case.id": "case_id",
                    "case.name": "case_name",
                    "case.data_set": "data_set",
                }.get(error.field, error.field)
                _location(error, workbook_path, "cases", row["_row"], column, row.get(column), "case contract")
                error.easytest_location["case_id"] = case_id
        raise
    for case in document["cases"]:
        case["steps"].sort(key=lambda item: int(item["order"]))
    return document


def compiled_content(path: str | Path) -> bytes:
    document = workbook_document(path)
    return (
        json.dumps(document, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    ).encode()


def compile_workbook(
    path: str | Path, output: str | Path | None = None, *, check: bool = False,
) -> Path:
    source = Path(path).resolve()
    target = Path(output).resolve() if output else source.with_suffix(".json")
    expected_target = source.with_suffix(".json")
    if target != expected_target:
        raise ContractError(
            f"compiled JSON must be the same-name sibling: {expected_target}",
            code="INVALID_OUTPUT_PATH", field="output",
        )
    content = compiled_content(source)
    if check:
        try:
            actual = target.read_bytes()
        except FileNotFoundError as exc:
            raise ContractError(
                f"compiled JSON is missing: {target}; run easytest compile {source}",
                code="COMPILED_JSON_MISSING", field="source",
            ) from exc
        if actual != content:
            raise ContractError(
                f"compiled JSON is out of date: {target}; run easytest compile {source}",
                code="COMPILED_JSON_OUT_OF_DATE", field="source",
            )
        return target
    if not target.exists() or target.read_bytes() != content:
        _atomic_write(target, content)
    return target


def compile_path(path: str | Path, *, check: bool = False) -> list[Path]:
    source = Path(path).resolve()
    candidates = [source] if source.is_file() else sorted(source.rglob("*"))
    workbooks = [
        path for path in candidates
        if path.is_file() and path.suffix.lower() == ".xlsx" and not path.name.startswith("~$")
    ]
    if not workbooks and (not check or not source.is_dir()):
        raise ContractError(f"no .xlsx case files found under {source}")
    outputs = [compile_workbook(workbook, check=check) for workbook in workbooks]
    if check and source.is_dir():
        paired = {path.resolve() for path in outputs}
        json_sources = 0
        for candidate in candidates:
            if (
                not candidate.is_file() or candidate.suffix.lower() != ".json"
                or candidate.resolve() in paired
            ):
                continue
            try:
                document = json.loads(candidate.read_text(encoding="utf-8"))
            except (UnicodeError, json.JSONDecodeError):
                continue
            if not isinstance(document, dict) or (
                "cases" not in document and "source_mode" not in document
            ):
                continue
            if document.get("source_mode") != "json":
                raise ContractError(
                    f"unpaired JSON must explicitly use source_mode='json': {candidate}",
                    code="ORPHAN_COMPILED_JSON", field="source_mode",
                )
            parse_document(document, source=str(candidate))
            json_sources += 1
        if not workbooks and not json_sources:
            raise ContractError(f"no XLSX or JSON case files found under {source}")
    return outputs
