from __future__ import annotations

import copy
import hashlib
import json
import os
import stat
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from easytest.cases.compiler import (
    CASE_COLUMNS,
    CASE_OPTIONAL_COLUMNS,
    JSON_COLUMNS,
    STEP_COLUMNS,
    STEP_OPTIONAL_COLUMNS,
    _atomic_write,
    compile_workbook,
    workbook_document,
)
from easytest.models import ContractError
from easytest.cases.schema import parse_document
from easytest.config import ProjectConfig
from easytest.runtime.policy import load_execution_policy
from easytest.runtime.preflight import preflight


_OPERATION_FIELDS = {
    "action",
    "entity",
    "case_id",
    "step_id",
    "new_id",
    "values",
}
_ACTIONS = {"add", "update", "rename", "delete"}
_ENTITIES = {"case", "step"}
_CASE_VALUES = (CASE_COLUMNS | CASE_OPTIONAL_COLUMNS) - {"case_id"}
_STEP_VALUES = (STEP_COLUMNS | STEP_OPTIONAL_COLUMNS) - {"case_id", "step_id"}


def _error(message: str, field_name: str = "edit") -> ContractError:
    return ContractError(message, code="INVALID_EDIT", field=field_name)


def _headers(worksheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in worksheet[1]
        if cell.value not in (None, "")
    }


def _ensure_columns(worksheet, names: set[str]) -> dict[str, int]:
    headers = _headers(worksheet)
    for name in sorted(names - headers.keys()):
        column = worksheet.max_column + 1
        cell = worksheet.cell(1, column, name)
        if column > 1:
            previous = worksheet.cell(1, column - 1)
            cell._style = copy.copy(previous._style)
            cell.font = copy.copy(previous.font)
            cell.fill = copy.copy(previous.fill)
            cell.border = copy.copy(previous.border)
            cell.alignment = copy.copy(previous.alignment)
            cell.number_format = previous.number_format
            cell.protection = copy.copy(previous.protection)
            cell.comment = copy.copy(previous.comment)
            previous_width = worksheet.column_dimensions[previous.column_letter].width
            worksheet.column_dimensions[cell.column_letter].width = previous_width
        headers[name] = column
    return headers


def _rows(worksheet, headers: dict[str, int], **keys: str) -> list[int]:
    found = []
    for row in range(2, worksheet.max_row + 1):
        if all(worksheet.cell(row, headers[name]).value == value for name, value in keys.items()):
            found.append(row)
    return found


def _cell_value(name: str, value: Any) -> Any:
    if name in JSON_COLUMNS and value is not None:
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    if name == "tags" and isinstance(value, list):
        if not all(isinstance(item, str) for item in value):
            raise _error("case tags must be strings", "values.tags")
        return ",".join(value)
    return value


def _set_values(worksheet, row: int, headers: dict[str, int], values: dict[str, Any]) -> None:
    for name, value in values.items():
        worksheet.cell(row, headers[name], _cell_value(name, value))


def _append_row(worksheet, headers: dict[str, int], values: dict[str, Any]) -> int:
    row = worksheet.max_row + 1
    for column in range(1, worksheet.max_column + 1):
        if worksheet.max_row >= 2:
            template = worksheet.cell(2, column)
            target = worksheet.cell(row, column)
            target._style = copy.copy(template._style)
            target.alignment = copy.copy(template.alignment)
            target.number_format = template.number_format
    _set_values(worksheet, row, headers, values)
    return row


def _validate_operation(value: Any, index: int) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise _error(f"edit operation {index} must be an object", f"operations[{index}]")
    unknown = value.keys() - _OPERATION_FIELDS
    if unknown:
        raise _error(
            f"edit operation {index} contains unknown fields: {sorted(unknown)}",
            f"operations[{index}].{sorted(unknown)[0]}",
        )
    action = value.get("action")
    entity = value.get("entity")
    if action not in _ACTIONS:
        raise _error(f"edit operation {index} has invalid action", f"operations[{index}].action")
    if entity not in _ENTITIES:
        raise _error(f"edit operation {index} has invalid entity", f"operations[{index}].entity")
    case_id = value.get("case_id")
    if not isinstance(case_id, str) or not case_id.strip():
        raise _error(f"edit operation {index} requires case_id", f"operations[{index}].case_id")
    step_id = value.get("step_id")
    if entity == "step" and (not isinstance(step_id, str) or not step_id.strip()):
        raise _error(f"edit operation {index} requires step_id", f"operations[{index}].step_id")
    if entity == "case" and step_id is not None:
        raise _error(f"edit operation {index} cannot use step_id", f"operations[{index}].step_id")
    values = value.get("values", {})
    if not isinstance(values, dict):
        raise _error(f"edit operation {index}.values must be an object", f"operations[{index}].values")
    allowed = _CASE_VALUES if entity == "case" else _STEP_VALUES
    invalid = values.keys() - allowed
    if invalid:
        raise _error(
            f"edit operation {index} contains unsupported values: {sorted(invalid)}",
            f"operations[{index}].values.{sorted(invalid)[0]}",
        )
    if action in {"add", "update"} and not values:
        raise _error(f"edit operation {index} requires values", f"operations[{index}].values")
    if action in {"rename", "delete"} and values:
        raise _error(f"edit operation {index} cannot combine {action} with values")
    new_id = value.get("new_id")
    if action == "rename" and (not isinstance(new_id, str) or not new_id.strip()):
        raise _error(f"edit operation {index} requires new_id", f"operations[{index}].new_id")
    if action != "rename" and new_id is not None:
        raise _error(f"edit operation {index} cannot use new_id", f"operations[{index}].new_id")
    return {
        "action": action,
        "entity": entity,
        "case_id": case_id.strip(),
        "step_id": step_id.strip() if isinstance(step_id, str) else None,
        "new_id": new_id.strip() if isinstance(new_id, str) else None,
        "values": values,
    }


def _apply_case(workbook, operation: dict[str, Any]) -> None:
    cases = workbook["cases"]
    steps = workbook["steps"]
    case_headers = _headers(cases)
    step_headers = _headers(steps)
    case_id = operation["case_id"]
    matches = _rows(cases, case_headers, case_id=case_id)
    action = operation["action"]

    if action == "add":
        if matches:
            raise _error(f"case already exists: {case_id}", "case_id")
        values = {"case_id": case_id, "case_name": case_id, **operation["values"]}
        missing = {"case_type"} - values.keys()
        if missing:
            raise _error(f"adding a case requires values: {sorted(missing)}", "values")
        case_headers = _ensure_columns(cases, set(values))
        _append_row(cases, case_headers, values)
        return
    if not matches:
        raise _error(f"case does not exist: {case_id}", "case_id")
    row = matches[0]
    if action == "update":
        case_headers = _ensure_columns(cases, set(operation["values"]))
        _set_values(cases, row, case_headers, operation["values"])
    elif action == "rename":
        new_id = operation["new_id"]
        if _rows(cases, case_headers, case_id=new_id):
            raise _error(f"case already exists: {new_id}", "new_id")
        cases.cell(row, case_headers["case_id"], new_id)
        for step_row in _rows(steps, step_headers, case_id=case_id):
            steps.cell(step_row, step_headers["case_id"], new_id)
    else:
        cases.delete_rows(row)
        for step_row in reversed(_rows(steps, step_headers, case_id=case_id)):
            steps.delete_rows(step_row)


def _apply_step(workbook, operation: dict[str, Any]) -> None:
    cases = workbook["cases"]
    steps = workbook["steps"]
    case_headers = _headers(cases)
    step_headers = _headers(steps)
    case_id = operation["case_id"]
    step_id = operation["step_id"]
    if not _rows(cases, case_headers, case_id=case_id):
        raise _error(f"case does not exist: {case_id}", "case_id")
    matches = _rows(steps, step_headers, case_id=case_id, step_id=step_id)
    action = operation["action"]

    if action == "add":
        if matches:
            raise _error(f"step already exists: {case_id}/{step_id}", "step_id")
        values = {"case_id": case_id, "step_id": step_id, **operation["values"]}
        missing = {"order", "executor", "operation"} - values.keys()
        if missing:
            raise _error(f"adding a step requires values: {sorted(missing)}", "values")
        step_headers = _ensure_columns(steps, set(values))
        _append_row(steps, step_headers, values)
        return
    if not matches:
        raise _error(f"step does not exist: {case_id}/{step_id}", "step_id")
    row = matches[0]
    if action == "update":
        step_headers = _ensure_columns(steps, set(operation["values"]))
        _set_values(steps, row, step_headers, operation["values"])
    elif action == "rename":
        new_id = operation["new_id"]
        if _rows(steps, step_headers, case_id=case_id, step_id=new_id):
            raise _error(f"step already exists: {case_id}/{new_id}", "new_id")
        steps.cell(row, step_headers["step_id"], new_id)
    else:
        steps.delete_rows(row)


def edit_workbook(
    path: str | Path,
    operations: list[dict[str, Any]],
    *,
    validate: bool = False,
    root: str | Path = ".",
    profile: str | None = None,
    execution_policy: str | Path | None = None,
) -> dict[str, Any]:
    if not validate and (profile is not None or execution_policy is not None):
        raise _error("profile/execution_policy require validate=True", "validate")
    source = Path(path).resolve()
    if source.suffix.lower() != ".xlsx":
        raise _error("structured editing only supports XLSX sources", "source")
    if not source.is_file():
        raise _error(f"workbook does not exist: {source}", "source")
    if not isinstance(operations, list) or not operations:
        raise _error("edit requires at least one operation", "operations")

    original = source.read_bytes()
    workbook_document(source)
    validated = [_validate_operation(item, index) for index, item in enumerate(operations)]
    temporary = None
    workbook = load_workbook(source)
    try:
        with tempfile.NamedTemporaryFile(
            dir=source.parent,
            prefix=f".{source.stem}.edit-",
            suffix=".xlsx",
            delete=False,
        ) as stream:
            temporary = Path(stream.name)
        for operation in validated:
            if operation["entity"] == "case":
                _apply_case(workbook, operation)
            else:
                _apply_step(workbook, operation)
        for worksheet in (workbook["cases"], workbook["steps"]):
            worksheet.auto_filter.ref = worksheet.dimensions
        workbook.save(temporary)
    except BaseException:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()

    try:
        document = workbook_document(temporary)
        if validate:
            config = ProjectConfig(root)
            settings = config.resolve_run(profile=profile)
            cases = [case for case in parse_document(document, source=str(source)) if case.enabled]
            if not cases:
                raise _error("edited workbook has no enabled cases", "enabled")
            preflight(
                cases, config, settings,
                execution_policy=load_execution_policy(root, execution_policy),
            )
        if source.read_bytes() != original:
            raise ContractError(
                f"workbook changed while editing: {source}",
                code="EDIT_CONFLICT",
                field="source",
            )
        os.chmod(temporary, stat.S_IMODE(source.stat().st_mode))
        os.replace(temporary, source)
        temporary = None
        try:
            compiled = compile_workbook(source)
        except BaseException:
            _atomic_write(source, original)
            raise
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)

    return {
        "source": str(source),
        "compiled": str(compiled),
        "operation_count": len(validated),
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "validated": validate,
    }
