from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from easytest.cases.compiler import compile_workbook
from easytest.models import ConfigurationError

_README = """# 我的测试项目

日常编辑 `cases/demo.xlsx`，运行 `easytest run` 或 `pytest`。
两个入口都会自动编译 Excel。项目默认离线运行，无需配置账号或启动服务。
AI 操作先读 `AI_GUIDE.md`，其中提供当前安装包的契约读取方法。
使用 `easytest list` 发现用例，`easytest run --case-id http.demo` 精确运行。
脚本化修改使用 `easytest edit cases/demo.xlsx --patch edit.json --validate --root .`，
落盘前预检工作簿；不要改生成的 JSON。跨文件校验使用 `easytest validate cases`。

如果在 EasyTest 框架目录创建了本项目，可以直接在那个环境执行：

```bash
uv run easytest run --root /本项目的绝对路径
uv run pytest /本项目的绝对路径/test_cases.py --easytest-root /本项目的绝对路径
```

也可以在本目录创建独立环境（EasyTest 尚未发布至 PyPI）：

```bash
uv venv
uv pip install /本机/easy_suite/easytest
source .venv/bin/activate
easytest run
pytest
```

Excel 的 cases 页维护用例名称，steps 页维护 operation、request、expect；
data 页集中维护可复用的多组业务数据。`sample.login` 是默认禁用的数据驱动样例，
启用后会按 data 页中 `login_cases` 的每个启用行独立执行。
技术列已预填并隐藏；新增用例/步骤时取消隐藏，复制行并更新唯一 ID 与 order。
`expect` 常用格式：`{"$.status_code": 200, "$.body.ok": true}`。
步骤可通过 `${data.username}` 引用当前数据行，完整占位符保留数字、布尔和 null 类型。
生成的同名 JSON 必须纳入版本控制；修改用例只编辑 Excel，再执行
`easytest compile cases/demo.xlsx`。提交前或 CI 使用
`easytest compile cases --check` 只读确认 JSON 与 XLSX 完全同步。
运行失败会给出 Case、Step、Operation、Excel 文件和 steps 页行号。

运行后直接打开 `artifacts/report.html` 查看统计、用例和步骤详情。
快照 diff 会自动按字段增删、类型、数值、文本、顺序等变化归类，
并汇总相同字段影响了多少用例。可搜索、筛选及下载报告数据。
报告可离线打开，无需启动服务；默认覆盖上次报告。
保留本次结果可使用 `easytest run --report artifacts/本次回归.html`，
或 `pytest --easytest-report artifacts/本次回归.html`。
机器读取可使用 `easytest run --result artifacts/result.json`，文件与 stdout 内容相同。
run / validate 成功、失败均输出 schema_version=1 的 JSON 封装：
status 表示状态，data 包含报告或预检摘要，errors 包含安全错误；日志和事件在 stderr。

真实接口测试：在 `config/operations.json` 配置 URL，执行
`easytest validate --profile live` 做只读预检，再执行
`easytest run --profile live`。运行时也会先检查整批用例。
validate 不生成 JSON、快照或报告，不调用业务服务或导入 handler；
结果中的 data.deferred 表示动态输出、handler 或截图等仍需运行时检查。
HTTP 顶层参数、Excel 列名和断言操作符均严格校验：
json 误写为 jsno、equals 误写为 equal 都会报错，业务 json 内部字段保持自由。
live 运行会真实访问所配置的服务。
`offline-strict` 禁止所有真实 Executor 调用；`mock=false` 或 inject 也不能绕过。
受控自动化可由宿主设置绝对路径的 `EASYTEST_EXECUTION_POLICY`，限制 Profile、
Case、operation、executor 及 HTTP method/origin，命令行不能覆盖该策略。
敏感配置使用环境变量，例如 `"Authorization": "Bearer ${HTTP_TOKEN}"`。
敏感 HTTP header 不接受配置字面量。HTTP 默认限制解压后响应体为 10 MiB，
超限显式失败，stream 不能绕过；详见安装包 CONTRACT.md。
项目可在 runtime.redaction 添加 secret_keys/pii_keys，但不能识别任意业务密钥。
`easytest run --fail-fast` 首个失败 Case 后停止，剩余项标为 not_run；不回滚业务写入。
项目根目录 .env 保存在独立映射中，进程环境优先；修改后重建 Runner / NotebookSession。

只有 `config/operations.json` 是必需文件。其余配置用于此项目的离线示例：
runtime 设置默认 Profile 与数据库写权限，profiles 定义运行预设，
mock_profiles 提供离线响应。没有开启快照，所以不需要快照配置或基线。

数据库写入默认关闭。确需造数时，显式使用 `--allow-db-write`；
快照 `--run-mode read/write/baseline` 独立控制比对、补齐和更新基线。
SQLite 快照的截图保存在 `snapshot_artifact_dir`，归档时必须与数据库一起处理。
使用 `easytest snapshot check --root .` 检查，维护前先停止其他快照写入。
Notebook 可以使用 `NotebookSession(".", profile="offline-strict")` 运行相同用例。
"""

_AI_GUIDE = """
先确认本项目根目录、接口契约、目标 Profile、允许修改/执行的 Case ID 和数据回收方式。
保留已有修改，不用放宽断言、更新快照或切换 Mock 掩盖失败。
已授权的操作直接完成；授权范围之外的业务决策再向负责人核实。

本项目以 `cases/demo.xlsx` 为唯一权威源，同名 JSON 是用于 Git diff、代码审查、
AI 分析和执行的确定性编译产物；两者一起提交，但只编辑 XLSX。
重命名/删除 XLSX 同步处理旧 JSON，否则会报孤立产物错误。
按 Case/Step ID 修改，隐藏列同样要维护唯一 ID 和 order。配置不要覆盖无关条目。
重复业务流程优先在 data 页增加数据行，并由 cases.data_set 引用；
步骤通过 `${data.字段名}` 读取当前行，不复制整套 Case/Step。
优先使用 `easytest edit cases/demo.xlsx --patch edit.json --validate --root .`
执行结构化增删改名；落盘前按项目配置预检工作簿，失败保留原文件，成功更新同名 JSON。
不加 --validate 仅检查工作簿契约；跨文件关系仍用 validate cases 检查。

```bash
easytest list --root .                      # 只读发现，无凭据解析或业务调用
easytest validate --root . --case-id http.demo --profile offline-strict
easytest compile cases/demo.xlsx            # 更新版本库中的同名 JSON
easytest compile cases --check              # 提交前/CI 只读检查同步
easytest run --root . --case-id http.demo --profile offline-strict --result artifacts/result.json
```

`--case-id` 可重复，精确匹配，未知/禁用 ID 或空选择失败；按来源顺序执行。
validate 不生成编译 JSON、存储、报告或业务调用；valid 只说明静态检查通过，
data.deferred 表示仍需运行检查。来源契约先整体校验，再预检选中 Case。
run 执行前也会整批预检，执行时会生成 XLSX 的编译产物。
生成 JSON 记录 XLSX 哈希和编译器版本；不要手工修改。JSON-only 仅用于没有
人工表格维护需求的纯代码/机器生成项目，必须显式设置 `source_mode: "json"`。

三个命令输出 schema_version=1 的 JSON。先看退出码、status，再读 data。
list=listed，validate=valid，run=passed；失败/中断非零退出。
run 的 data.summary/cases 是报告，errors 有 code/field 和可用的来源定位。
validate.data.input_hash 与 run.data.input_hash 可核对所选 Case、配置、运行设置和
执行策略是否一致；该哈希不包含环境变量值、handler 源码或外部状态。
普通断言 errors 含 path/operator/expected/actual，敏感值已脱敏。
只有 passed 步骤的 mocked 才能用于区分 Mock 与真实执行。
日志在 stderr，--result 与 stdout 相同，失败也可能生成 artifacts/report.html。

live 允许真实执行但不会清除显式 Mock；read 只控制快照，不能阻止 HTTP/RPC 写入。
数据库写权限单独控制。失败后的业务状态可能未知，不要盲目重跑写操作。
CLI 可使用 --fail-fast 或 --max-failures N，后续未执行 Case 为 not_run。
HTTP_RESPONSE_TOO_LARGE 的响应哈希仅覆盖已读取前缀，不代表完整响应；
该步骤未执行断言或快照，也不自动重试。先核对接口与副作用再处理。
SNAPSHOT_CONFLICT 表示同一 Case 已被其他进程更新；重新读取基线，不直接覆盖。
SQLite 图片目录和数据库共同组成快照，维护使用 `easytest snapshot check/maintain`。
真实自动化应由可信宿主设置 `EASYTEST_EXECUTION_POLICY`；它是应用层 allowlist，
不是操作系统网络沙箱。受保护的 HTTP operation 需设置 `allow_redirects=false`；
获准的自定义 handler 仍属于受信任代码。
清理用业务 finally/pytest fixture，不依赖最后一个步骤。凭据留在本地 .env 或 CI。

从当前 Python 环境读取版本和随包契约，不依赖框架源码目录：

```bash
python -c "import easytest; print(easytest.__version__, easytest.__file__)"
python -c "from importlib.resources import files; print(files('easytest').joinpath('CONTRACT.md').read_text(encoding='utf-8'))"
```

交付说明修改的源文件/ID、实际命令、静态/Mock/真实结果、失败位置、
未执行项、已发生的业务动作和报告路径。本说明不是权限或网络沙箱。
"""


def init_project(directory: str | Path) -> Path:
    """Create only user-facing cases, configuration and one pytest entry."""
    from easytest import __version__

    root = Path(directory).absolute()
    if root.is_symlink() or (
        root.exists() and (not root.is_dir() or any(root.iterdir()))
    ):
        raise ConfigurationError(f"init requires a new or empty directory: {root}")
    root.mkdir(parents=True, exist_ok=True)
    config = root / "config"
    config.mkdir()
    documents = {
        "operations": {"operations": {
            "http.ping": {"executor": "http", "method": "GET", "url": "https://example.invalid/ping"},
            "http.login": {"executor": "http", "method": "POST", "url": "https://example.invalid/login"},
        }},
        "runtime": {"default_profile": "offline-strict", "allow_db_write": False},
        "profiles": {"profiles": {
            "offline-strict": {"mock_profile": "offline", "fail_on_mock_miss": True},
            "live": {"fail_on_mock_miss": False},
        }},
        "mock_profiles": {"profiles": {"offline": {
            "http.ping": {"kind": "response", "response": {
                "status_code": 200, "body": {"ok": True},
            }},
        }}},
    }
    for name, document in documents.items():
        (config / f"{name}.json").write_text(
            json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8",
        )
    (root / "README.md").write_text(_README, encoding="utf-8")
    (root / "AI_GUIDE.md").write_text(
        f"# EasyTest {__version__} AI 接入简明指南\n{_AI_GUIDE}", encoding="utf-8",
    )
    (root / "pytest.ini").write_text(
        "[pytest]\ntestpaths = test_cases.py\naddopts = -ra\n", encoding="utf-8",
    )
    (root / "test_cases.py").write_text(
        "def test_case(table_case, case_runner):\n    case_runner.run(table_case)\n",
        encoding="utf-8",
    )
    (root / ".gitignore").write_text(
        ".venv/\n.env\n__pycache__/\n.pytest_cache/\n.easytest/\nartifacts/\n~$*.xlsx\n",
        encoding="utf-8",
    )
    workbook = Workbook()
    cases = workbook.active
    cases.title = "cases"
    cases.append(["case_id", "case_name", "case_type", "enabled", "data_set"])
    cases.append(["http.demo", "检查接口返回成功", "http", True, None])
    cases.append([
        "sample.login",
        "登录数据驱动示例（启用后运行）",
        "http",
        False,
        "login_cases",
    ])
    steps = workbook.create_sheet("steps")
    steps.append([
        "case_id",
        "step_id",
        "order",
        "executor",
        "operation",
        "request",
        "mock",
        "expect",
    ])
    steps.append([
        "http.demo", "ping", 1, "http", "http.ping", "{}",
        None,
        '{"$.status_code": 200, "$.body.ok": true}',
    ])
    steps.append([
        "sample.login",
        "login",
        1,
        "http",
        "http.login",
        '{"json":{"username":"${data.username}",'
        '"password":"${data.password}","phone":"${data.phone}"}}',
        '{"response":{"status_code":"${data.expected_status}",'
        '"body":{"message":"${data.expected_message}"}}}',
        '{"$.status_code":"${data.expected_status}",'
        '"$.body.message":"${data.expected_message}"}',
    ])
    data = workbook.create_sheet("data")
    data.append([
        "data_set",
        "data_id",
        "enabled",
        "username",
        "password",
        "phone",
        "expected_status",
        "expected_message",
    ])
    data.append([
        "login_cases", "valid", True, "demo-user", "correct-demo-only",
        "13800000000", 200, "login succeeded",
    ])
    data.append([
        "login_cases", "wrong_password", True, "demo-user", "wrong-demo-only",
        "13800000000", 401, "invalid credentials",
    ])
    data.append([
        "login_cases", "empty_phone", True, "demo-user", "correct-demo-only",
        None, 400, "phone is required",
    ])
    for sheet, hidden in (
        (cases, ("A", "C")),
        (steps, ("A", "B", "C", "D")),
        (data, ("A",)),
    ):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.comment = Comment("已预填；新增行时注意唯一 ID 和执行顺序。", "EasyTest")
            sheet.column_dimensions[cell.column_letter].width = 36
        for column in hidden:
            sheet.column_dimensions[column].hidden = True
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[cell.row].height = 45
    source = root / "cases" / "demo.xlsx"
    source.parent.mkdir()
    try:
        workbook.save(source)
    finally:
        workbook.close()
    compile_workbook(source)
    return root
