from __future__ import annotations

import hashlib
import json
from pathlib import Path
import tomllib

import pytest
from openpyxl import Workbook, load_workbook

from easytest import __version__
from easytest.cases.compiler import compile_path, compile_workbook
from easytest.cases.loader import load_cases
from easytest.cli import main
from easytest.models import ContractError


def _minimal_workbook(path: Path, *, formula: bool = False) -> None:
    workbook = Workbook()
    cases = workbook.active
    cases.title = "cases"
    cases.append(
        [
            "case_id",
            "case_name",
            "case_type",
            "enabled",
            "variables",
            "mock_profile",
            "snapshot_profile",
        ]
    )
    cases.append(
        [
            "case.one",
            '=UPPER("name")' if formula else "Case One",
            "http",
            True,
            "{}",
            "offline",
            "default",
        ]
    )
    steps = workbook.create_sheet("steps")
    steps.append(
        [
            "case_id",
            "step_id",
            "order",
            "executor",
            "operation",
            "request",
        ]
    )
    steps.append(["case.one", "request", 1, "http", "http.get_profile", "{}"])
    workbook.save(path)


def test_compile_workbook_is_deterministic(tmp_path: Path) -> None:
    source = tmp_path / "cases.xlsx"
    _minimal_workbook(source)

    output = compile_workbook(source)
    first = output.read_bytes()
    compile_workbook(source)

    assert output.read_bytes() == first
    document = json.loads(first)
    assert document["source_mode"] == "xlsx"
    assert document["source"] == "cases.xlsx"
    assert document["source_sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()
    assert document["compiler_version"] == __version__
    assert load_cases(output)[0].id == "case.one"
    assert load_cases(output)[0].source == str(source)
    assert load_cases(output)[0].steps[0].source_row == 2


def test_package_and_project_versions_match():
    project = tomllib.loads(
        (Path(__file__).resolve().parents[1] / "pyproject.toml").read_text()
    )
    assert project["project"]["version"] == __version__


def test_compile_rejects_non_sibling_output(tmp_path):
    source = tmp_path / "cases.xlsx"
    _minimal_workbook(source)
    with pytest.raises(ContractError) as caught:
        compile_workbook(source, tmp_path / "review.json")
    assert caught.value.code == "INVALID_OUTPUT_PATH"
    assert not (tmp_path / "review.json").exists()


def test_compile_check_is_read_only_and_detects_missing_or_changed_output(tmp_path):
    source = tmp_path / "cases.xlsx"
    _minimal_workbook(source)
    output = compile_workbook(source)
    before = {
        path.name: (path.read_bytes(), path.stat().st_mtime_ns)
        for path in tmp_path.iterdir()
    }
    assert compile_path(tmp_path, check=True) == [output]
    assert {
        path.name: (path.read_bytes(), path.stat().st_mtime_ns)
        for path in tmp_path.iterdir()
    } == before

    output.write_text("{}\n")
    changed = output.read_bytes()
    with pytest.raises(ContractError) as caught:
        compile_workbook(source, check=True)
    assert caught.value.code == "COMPILED_JSON_OUT_OF_DATE"
    assert output.read_bytes() == changed

    output.unlink()
    with pytest.raises(ContractError) as caught:
        compile_workbook(source, check=True)
    assert caught.value.code == "COMPILED_JSON_MISSING"
    assert not output.exists()


def test_cli_compile_check(tmp_path, capsys):
    source = tmp_path / "cases.xlsx"
    _minimal_workbook(source)
    output = compile_workbook(source)
    main(["compile", str(source), "--check"])
    assert capsys.readouterr().out.strip() == str(output)
    output.write_text("{}\n")
    with pytest.raises(ContractError) as caught:
        main(["compile", str(source), "--check"])
    assert caught.value.code == "COMPILED_JSON_OUT_OF_DATE"


def test_compiled_json_detects_changed_workbook_before_direct_loading(tmp_path):
    source = tmp_path / "cases.xlsx"
    _minimal_workbook(source)
    output = compile_workbook(source)
    workbook = load_workbook(source)
    workbook["cases"]["B2"] = "Changed"
    workbook.save(source)
    workbook.close()
    with pytest.raises(ContractError) as caught:
        load_cases(output)
    assert caught.value.code == "COMPILED_JSON_OUT_OF_DATE"


def test_generated_and_json_only_modes_are_explicit(tmp_path):
    source = tmp_path / "cases.xlsx"
    _minimal_workbook(source)
    output = compile_workbook(source)
    document = json.loads(output.read_text())

    document.pop("source_mode")
    output.write_text(json.dumps(document))
    with pytest.raises(ContractError) as caught:
        load_cases(output)
    assert caught.value.code == "COMPILED_JSON_OUT_OF_DATE"

    document["source_mode"] = "json"
    output.write_text(json.dumps(document))
    with pytest.raises(ContractError) as caught:
        load_cases(output)
    assert caught.value.code == "INVALID_SOURCE_METADATA"


def test_directory_check_rejects_orphans_and_accepts_explicit_json_only(tmp_path):
    source = tmp_path / "cases.xlsx"
    _minimal_workbook(source)
    output = compile_workbook(source)
    source.unlink()
    with pytest.raises(ContractError) as caught:
        compile_path(tmp_path, check=True)
    assert caught.value.code == "ORPHAN_COMPILED_JSON"

    document = json.loads(output.read_text())
    for field in ("source", "source_sha256", "compiler_version"):
        document.pop(field)
    document["source_mode"] = "json"
    output.write_text(json.dumps(document))
    before = output.read_bytes()
    assert compile_path(tmp_path, check=True) == []
    assert output.read_bytes() == before


def test_compile_rejects_formula_cells(tmp_path: Path) -> None:
    source = tmp_path / "cases.xlsx"
    _minimal_workbook(source, formula=True)

    with pytest.raises(ContractError, match="contains a formula"):
        compile_workbook(source)


def test_runtime_rejects_xlsx_input(tmp_path: Path) -> None:
    source = tmp_path / "cases.xlsx"
    _minimal_workbook(source)

    with pytest.raises(ContractError, match="only accepts compiled .json"):
        load_cases(source)


def test_invalid_step_reports_excel_sheet_and_row(tmp_path):
    source = tmp_path / "bad.xlsx"
    _minimal_workbook(source)
    workbook = load_workbook(source)
    workbook["steps"]["C2"] = "not-an-order"
    workbook.save(source)
    workbook.close()
    with pytest.raises(ContractError, match="order must be an integer") as error:
        compile_workbook(source)
    assert "sheet=steps, row=2" in str(error.value.__notes__)
    assert str(source) in str(error.value.__notes__)
