from __future__ import annotations

import json
import sqlite3
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from easytest.cases.compiler import compile_workbook, workbook_document
from easytest.cases.loader import load_cases, load_project_cases
from easytest.executors.base import Executor
from easytest.models import ContractError, ExecutionResult
from easytest.reports.results import RunReport
from easytest.runtime.runner import CaseRunner
from easytest.starter import init_project


def _write_data_workbook(
    path: Path,
    *,
    executor: str = "http",
    operation: str = "http.ping",
    request: dict | None = None,
    mock: dict | bool | None = None,
    expect: dict | bool | None = None,
    rows: list[list[object]] | None = None,
) -> None:
    workbook = Workbook()
    cases = workbook.active
    cases.title = "cases"
    cases.append(
        [
            "case_id",
            "case_name",
            "case_type",
            "enabled",
            "variables",
            "data_set",
        ]
    )
    cases.append(
        [
            "login.case",
            "登录数据驱动",
            "http",
            True,
            '{"tenant":"qa"}',
            "login_cases",
        ]
    )

    steps = workbook.create_sheet("steps")
    steps.append(
        [
            "case_id",
            "step_id",
            "order",
            "executor",
            "operation",
            "request",
            "mock",
            "expect",
        ]
    )
    steps.append(
        [
            "login.case",
            "login",
            1,
            executor,
            operation,
            json.dumps(
                request
                or {
                    "json": {
                        "username": "${data.username}",
                        "password": "${data.password}",
                        "phone": "${data.phone}",
                        "tenant": "${variables.tenant}",
                    }
                },
                ensure_ascii=False,
            ),
            json.dumps(
                mock
                if isinstance(mock, dict)
                else {
                    "response": {
                        "status_code": "${data.expected_status}",
                        "body": {"username": "${data.username}"},
                    },
                },
                ensure_ascii=False,
            ) if mock is not False else "",
            json.dumps(
                expect
                if isinstance(expect, dict)
                else {"$.status_code": "${data.expected_status}"},
                ensure_ascii=False,
            ) if expect is not False else "",
        ]
    )

    data = workbook.create_sheet("data")
    data.append(
        [
            "data_set",
            "data_id",
            "enabled",
            "username",
            "password",
            "phone",
            "expected_status",
        ]
    )
    for row in rows or [
        ["login_cases", "valid", True, "alice", "correct", "13800000000", 200],
        ["login_cases", "wrong_password", True, "alice", "wrong", "13800000000", 401],
        ["login_cases", "empty_phone", True, "bob", "correct", None, 400],
    ]:
        data.append(row)
    workbook.save(path)
    workbook.close()


def test_data_sheet_compiles_and_expands_enabled_rows(tmp_path: Path) -> None:
    source = tmp_path / "login.xlsx"
    _write_data_workbook(
        source,
        rows=[
            ["login_cases", "valid", True, "alice", "correct", "13800000000", 200],
            ["login_cases", "disabled", False, "bob", "wrong", None, 401],
            ["login_cases", "empty_phone", True, "carol", "correct", None, 400],
        ],
    )

    compiled = compile_workbook(source)
    document = json.loads(compiled.read_text(encoding="utf-8"))
    cases = load_cases(compiled)

    assert document["cases"][0]["data_set"] == "login_cases"
    assert [row["id"] for row in document["data_sets"]["login_cases"]] == [
        "valid",
        "disabled",
        "empty_phone",
    ]
    assert [case.id for case in cases] == ["login.case", "login.case"]
    assert [case.execution_id for case in cases] == [
        "login.case.valid",
        "login.case.empty_phone",
    ]
    assert [case.data_id for case in cases] == ["valid", "empty_phone"]
    assert cases[1].data == {
        "username": "carol",
        "password": "correct",
        "phone": None,
        "expected_status": 400,
    }
    assert cases[1].data_source_row == 4


def test_case_selection_accepts_logical_or_data_execution_id(
    tmp_path: Path,
) -> None:
    source = tmp_path / "login.xlsx"
    _write_data_workbook(source)

    all_rows = load_project_cases(
        source,
        root=tmp_path,
        write_compiled=False,
        case_ids=["login.case"],
    )
    one_row = load_project_cases(
        source,
        root=tmp_path,
        write_compiled=False,
        case_ids=["login.case.wrong_password"],
    )

    assert [case.data_id for case in all_rows] == [
        "valid",
        "wrong_password",
        "empty_phone",
    ]
    assert [case.data_id for case in one_row] == ["wrong_password"]


def test_init_includes_a_complete_disabled_data_driven_example(
    tmp_path: Path,
) -> None:
    root = init_project(tmp_path / "project")
    source = root / "cases/demo.xlsx"
    workbook = load_workbook(source, read_only=True)
    try:
        assert workbook.sheetnames == ["cases", "steps", "data"]
        assert workbook["cases"]["E3"].value == "login_cases"
        assert [workbook["data"].cell(row, 2).value for row in range(2, 5)] == [
            "valid",
            "wrong_password",
            "empty_phone",
        ]
    finally:
        workbook.close()

    document = json.loads(source.with_suffix(".json").read_text(encoding="utf-8"))
    example = next(case for case in document["cases"] if case["id"] == "sample.login")
    assert example["enabled"] is False
    assert example["data_set"] == "login_cases"
    assert len(document["data_sets"]["login_cases"]) == 3
    assert [case.id for case in load_project_cases(root / "cases")] == ["http.demo"]

    workbook = load_workbook(source)
    workbook["cases"]["D3"] = True
    workbook.save(source)
    workbook.close()
    compile_workbook(source)
    examples = load_project_cases(
        source,
        root=root,
        case_ids=["sample.login"],
    )
    with CaseRunner(root) as runner:
        for case in examples:
            runner.run(case)
    assert [case.data_id for case in examples] == [
        "valid",
        "wrong_password",
        "empty_phone",
    ]
    assert [report.status for report in runner.case_reports] == [
        "passed",
        "passed",
        "passed",
    ]


def test_each_data_row_runs_with_an_isolated_context_and_visible_report(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = init_project(tmp_path / "project")
    source = root / "cases/login.xlsx"
    _write_data_workbook(source, mock=False)
    compile_workbook(source)
    cases = load_project_cases(source, root=root)

    seen: list[tuple[str, dict, dict]] = []
    instances = []

    class CaptureHttp(Executor):
        def __init__(self) -> None:
            self.closed = 0
            self.resets = 0
            instances.append(self)

        def reset_context(self, _context) -> None:
            self.resets += 1

        def execute(self, operation_name, _operation, request, context):
            seen.append(
                (
                    context.case.execution_id,
                    dict(context.state),
                    request,
                )
            )
            context.state["visited"] = context.case.data_id
            return ExecutionResult(
                "http",
                operation_name,
                {
                    "status_code": context.case.data["expected_status"],
                    "body": {"username": context.case.data["username"]},
                },
            )

        def close(self) -> None:
            self.closed += 1

    monkeypatch.setattr("easytest.runtime.runner.HttpExecutor", CaptureHttp)
    with CaseRunner(root, profile="live") as runner:
        checked = runner.preflight(cases, reuse_for_run=True)
        for case in cases:
            runner.run(case)

        assert checked.case_count == 3
        assert checked.step_count == 3
        assert len(instances) == 1
        assert instances[0].resets == 2
        assert all(state == {} for _, state, _ in seen)
        assert [request["json"]["phone"] for _, _, request in seen] == [
            "13800000000",
            "13800000000",
            None,
        ]
        assert all(
            request["json"]["tenant"] == "qa"
            for _, _, request in seen
        )
        reports = RunReport(cases=runner.case_reports).as_dict()["cases"]
        assert [report["data_id"] for report in reports] == [
            "valid",
            "wrong_password",
            "empty_phone",
        ]
        assert reports[0]["data"]["password"] == "***REDACTED***"
        assert reports[0]["execution_id"] == "login.case.valid"

    assert instances[0].closed == 1


def test_data_values_bind_to_database_parameters(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = init_project(tmp_path / "project")
    database = root / "users.db"
    connection = sqlite3.connect(database)
    connection.execute("CREATE TABLE users (username TEXT, status TEXT)")
    connection.executemany(
        "INSERT INTO users VALUES (?, ?)",
        [("alice", "ACTIVE"), ("bob", "LOCKED")],
    )
    connection.commit()
    connection.close()

    runtime = json.loads((root / "config/runtime.json").read_text())
    runtime["database"] = {
        "connections": {
            "default": {
                "driver": "sqlite",
                "database": "users.db",
                "read_only": True,
            }
        }
    }
    (root / "config/runtime.json").write_text(json.dumps(runtime))
    operations = json.loads((root / "config/operations.json").read_text())
    operations["operations"]["database.user"] = {
        "executor": "database",
        "connection": "default",
        "statement": (
            "SELECT username, status FROM users "
            "WHERE username = :username"
        ),
        "write": False,
    }
    (root / "config/operations.json").write_text(json.dumps(operations))

    source = root / "cases/users.xlsx"
    _write_data_workbook(
        source,
        executor="database",
        operation="database.user",
        request={"parameters": {"username": "${data.username}"}},
        mock=False,
        expect=False,
        rows=[
            ["login_cases", "alice", True, "alice", "secret-a", None, 200],
            ["login_cases", "bob", True, "bob", "secret-b", None, 200],
        ],
    )
    cases = load_project_cases(source, root=root)

    import easytest.executors.database as database_module

    real_connect = database_module.sqlite3.connect
    connection_count = 0

    def tracked_connect(*args, **kwargs):
        nonlocal connection_count
        connection_count += 1
        return real_connect(*args, **kwargs)

    monkeypatch.setattr(database_module.sqlite3, "connect", tracked_connect)
    with CaseRunner(root, profile="live") as runner:
        outputs = [
            runner.run(case)["login"].output["rows"][0]
            for case in cases
        ]

    assert outputs == [
        {"username": "alice", "status": "ACTIVE"},
        {"username": "bob", "status": "LOCKED"},
    ]
    assert connection_count == 1


def test_data_contract_errors_report_sheet_row_and_column(tmp_path: Path) -> None:
    duplicate = tmp_path / "duplicate.xlsx"
    _write_data_workbook(
        duplicate,
        rows=[
            ["login_cases", "same", True, "alice", "one", None, 200],
            ["login_cases", "same", True, "bob", "two", None, 401],
        ],
    )
    with pytest.raises(ContractError) as caught:
        workbook_document(duplicate)
    assert caught.value.code == "DUPLICATE_DATA_ID"
    assert caught.value.easytest_location == {
        "source": str(duplicate),
        "sheet": "data",
        "source_row": 3,
        "column": "data_id",
        "data_set": "login_cases",
        "data_id": "same",
    }

    missing = tmp_path / "missing.xlsx"
    _write_data_workbook(missing)
    workbook = load_workbook(missing)
    workbook["cases"]["F2"] = "missing_set"
    workbook.save(missing)
    workbook.close()
    with pytest.raises(ContractError) as caught:
        workbook_document(missing)
    assert caught.value.code == "UNKNOWN_DATA_SET"
    assert caught.value.easytest_location["sheet"] == "cases"
    assert caught.value.easytest_location["source_row"] == 2
    assert caught.value.easytest_location["column"] == "data_set"

    disabled = tmp_path / "disabled.xlsx"
    _write_data_workbook(
        disabled,
        rows=[
            ["login_cases", "disabled", False, "alice", "one", None, 200],
        ],
    )
    with pytest.raises(ContractError) as caught:
        workbook_document(disabled)
    assert caught.value.code == "EMPTY_DATA_SET"


def test_data_rows_use_independent_snapshot_namespaces(tmp_path: Path) -> None:
    root = init_project(tmp_path / "project")
    (root / "config/snapshots.json").write_text(
        json.dumps(
            {"profiles": {"default": {"response_default": {}}}}
        ),
        encoding="utf-8",
    )
    source = root / "cases/login.xlsx"
    _write_data_workbook(
        source,
        rows=[
            ["login_cases", "valid", True, "alice", "correct", None, 200],
            ["login_cases", "wrong_password", True, "alice", "wrong", None, 401],
        ],
    )
    cases = [
        replace(
            case,
            steps=(
                replace(
                    case.steps[0],
                    snapshot={"rule": "response_default", "name": "login"},
                ),
            ),
        )
        for case in load_project_cases(source, root=root, write_compiled=False)
    ]

    with CaseRunner(root, run_mode="write") as runner:
        for case in cases:
            runner.run(case)

    assert (root / "snapshots/login.case.valid/login.json").is_file()
    assert (
        root / "snapshots/login.case.wrong_password/login.json"
    ).is_file()


def test_missing_data_field_fails_before_execution_with_data_location(
    tmp_path: Path,
) -> None:
    root = init_project(tmp_path / "project")
    source = root / "cases/login.xlsx"
    _write_data_workbook(
        source,
        request={"json": {"missing": "${data.not_defined}"}},
        rows=[
            ["login_cases", "valid", True, "alice", "correct", None, 200],
        ],
    )
    cases = load_project_cases(source, root=root, write_compiled=False)

    with CaseRunner(root) as runner:
        with pytest.raises(ContractError, match="does not exist") as caught:
            runner.preflight(cases)

    assert caught.value.preflight_data_id == "valid"
    assert caught.value.preflight_data_source_row == 2
    assert "data_id=valid" in str(caught.value.__notes__)
