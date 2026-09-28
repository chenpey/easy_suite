from datetime import datetime

import pytest
from openpyxl import load_workbook

from easytest.cases.compiler import workbook_document
from easytest.models import ContractError
from easytest.reports.results import error_info
from easytest.starter import init_project


@pytest.mark.parametrize(
    "column,value,actual_type",
    [("case_id", 123, "int"), ("case_id", datetime(2026, 1, 1), "datetime"),
     ("case_name", 123, "int"), ("case_type", "typo", "str"),
     ("variables", "[]", "list"), ("enabled", "not-a-bool", "str")],
)
def test_case_cell_errors_have_sheet_row_column_and_types(tmp_path, column, value, actual_type):
    root = init_project(tmp_path / "project")
    source = root / "cases/demo.xlsx"
    workbook = load_workbook(source)
    sheet = workbook["cases"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    if column not in headers:
        headers[column] = sheet.max_column + 1
        sheet.cell(1, headers[column], column)
    sheet.cell(2, headers[column], value)
    workbook.save(source)
    workbook.close()
    before = source.read_bytes()
    with pytest.raises(ContractError) as caught:
        workbook_document(source)
    result = error_info(caught.value, "collection")
    assert result["source"] == str(source)
    assert result["sheet"] == "cases"
    assert result["source_row"] == "2"
    assert result["column"] == column
    assert result["actual_type"] == actual_type
    assert result["expected_type"]
    if column == "case_id":
        assert "Format the Excel cell as Text" in str(caught.value)
        assert "leading zeroes" in str(caught.value)
    assert source.read_bytes() == before


def test_text_identifier_preserves_leading_zeroes(tmp_path):
    root = init_project(tmp_path / "project")
    source = root / "cases/demo.xlsx"
    workbook = load_workbook(source)
    workbook["cases"]["A2"] = "001"
    workbook["steps"]["A2"] = "001"
    workbook.save(source)
    workbook.close()
    assert workbook_document(source)["cases"][0]["id"] == "001"


def test_unknown_excel_column_suggests_nearest_header(tmp_path):
    root = init_project(tmp_path / "project")
    source = root / "cases/demo.xlsx"
    workbook = load_workbook(source)
    workbook["cases"]["A1"] = "caseid"
    workbook.save(source)
    workbook.close()

    with pytest.raises(ContractError, match="did you mean 'case_id'"):
        workbook_document(source)


def test_invalid_case_type_suggests_nearest_value(tmp_path):
    root = init_project(tmp_path / "project")
    source = root / "cases/demo.xlsx"
    workbook = load_workbook(source)
    workbook["cases"]["C2"] = "scenrio"
    workbook.save(source)
    workbook.close()

    with pytest.raises(ContractError, match="did you mean 'scenario'"):
        workbook_document(source)
