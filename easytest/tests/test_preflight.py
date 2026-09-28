from __future__ import annotations

import json
import subprocess
import sys
import threading
from dataclasses import replace
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest
from openpyxl import load_workbook

from easytest.cases.compiler import workbook_document
from easytest.cases.loader import load_cases, load_project_cases
from easytest.cli import main
from easytest.config import ProjectConfig
from easytest.executors.base import Executor
from easytest.models import ConfigurationError, ContractError, ExecutionResult
from easytest.notebook import NotebookSession
from easytest.runtime.assertions import assert_expectations
from easytest.runtime.preflight import preflight
from easytest.runtime.runner import CaseRunner
from easytest.snapshots.store import SqliteSnapshotStore
from easytest.starter import init_project


@pytest.fixture
def server():
    requests = []

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            body = self.rfile.read(int(self.headers.get("Content-Length", 0)))
            requests.append(json.loads(body) if body else None)
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(b'{"ok":true,"id":7,"bad_request":{"jsno":{}}}')

        def log_message(self, *_args):
            pass

    service = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=service.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{service.server_port}/items", requests
    finally:
        service.shutdown()
        service.server_close()
        thread.join()


def _project(tmp_path, url):
    root = init_project(tmp_path / "project")
    (root / "config/operations.json").write_text(json.dumps({"operations": {
        "http.ping": {"executor": "http", "url": url, "method": "POST", "trust_env": False},
    }}))
    return root


def _document(root):
    return json.loads((root / "cases/demo.json").read_text())


def _save(root, document):
    path = root / "cases/demo.json"
    path.write_text(json.dumps(document))
    return path


@pytest.mark.parametrize("change,error", [
    ({"operation": "http.missing"}, "unknown operation"),
    ({"executor": "rpc"}, "but step requested"),
    ({"request": {"json": "${variables.missing}"}}, "does not exist"),
    ({"request": {"json": "${steps.future.body}"}}, "does not exist"),
    ({"request": {"jsno": {"id": 1}}}, "unknown fields"),
    ({"expect": {"checks": [{"path": "$.status_code", "equal": 418}]}}, "unknown expectation"),
    ({"snapshot": "missing"}, "unknown snapshot rule"),
    ({"mock": "missing"}, "unknown mock preset"),
    ({"request": {"url": "https://example.invalid/items/{id}"}}, "path parameter"),
    ({"request": {"timeout": "soon"}}, "HTTP timeout"),
    ({"request": {"retry": {"retrise": 2}}}, "unknown fields"),
])
@pytest.mark.parametrize("across_cases", [False, True])
def test_later_static_error_prevents_all_http_requests(tmp_path, server, change, error, across_cases):
    url, received = server
    root = _project(tmp_path, url)
    document = _document(root)
    first = document["cases"][0]
    later = {**first["steps"][0], "id": "later", "order": 2, **change}
    if across_cases:
        document["cases"].append({**first, "id": "later.case", "steps": [later]})
    else:
        first["steps"].append(later)
    path = _save(root, document)
    with pytest.raises((ConfigurationError, ContractError), match=error):
        main(["run", str(path), "--root", str(root), "--profile", "live", "--no-report"])
    assert received == []


def test_single_case_preflight_keeps_failure_location_and_not_run_steps(tmp_path, server):
    url, received = server
    root = _project(tmp_path, url)
    case = load_cases(root / "cases/demo.json")[0]
    case = replace(case, steps=(
        case.steps[0], replace(case.steps[0], id="later", order=2, operation="missing"),
    ))
    with CaseRunner(root, profile="live") as runner:
        with pytest.raises(ConfigurationError) as caught:
            runner.run(case)
        assert runner.last_report.errors[0]["phase"] == "preflight"
        assert [step.status for step in runner.last_report.steps] == ["not_run", "failed"]
        assert "sheet=steps, row=2" in str(caught.value.__notes__)
    assert received == []


def test_pytest_checks_selected_batch_before_any_fixture_or_request(tmp_path, server):
    url, received = server
    root = _project(tmp_path, url)
    document = _document(root)
    first = document["cases"][0]
    document["cases"].append({
        **first, "id": "bad.case",
        "steps": [{**first["steps"][0], "operation": "missing"}],
    })
    _save(root, document)
    command = [
        sys.executable, "-m", "pytest", "-q", "--case-source", "cases/demo.json",
        "--profile", "live", "--no-easytest-report",
    ]
    blocked = subprocess.run(command, cwd=root, text=True, capture_output=True, timeout=30)
    assert blocked.returncode != 0
    assert "preflight failed" in blocked.stdout + blocked.stderr
    assert received == []
    # Deselecting the bad case must allow the remaining valid selection.
    selected = subprocess.run(
        [*command, "-k", "http.demo"], cwd=root, text=True, capture_output=True, timeout=30,
    )
    assert selected.returncode == 0, selected.stdout + selected.stderr
    assert len(received) == 1


def _tree(root):
    return {
        str(path.relative_to(root)): (path.read_bytes(), path.stat().st_mtime_ns)
        for path in root.rglob("*") if path.is_file()
    }


def test_validate_reads_excel_without_writing_json_snapshots_or_reports(tmp_path, capsys):
    root = init_project(tmp_path / "project")
    (root / "cases/demo.json").write_text("stale JSON must not be parsed")
    (root / "config/runtime.json").write_text(json.dumps({
        "default_profile": "offline-strict", "snapshot_backend": "sqlite",
    }))
    before = _tree(root)
    main(["validate", "--root", str(root)])
    result = json.loads(capsys.readouterr().out)
    assert result["status"] == "valid"
    assert result["data"] == {
        "status": "valid",
        "case_count": 1,
        "step_count": 1,
        "deferred": [],
        "input_hash": result["data"]["input_hash"],
    }
    assert result["data"]["input_hash"].startswith("sha256:")
    assert _tree(root) == before


def test_offline_needs_no_credentials_but_live_reports_only_missing_name(tmp_path, monkeypatch, capsys):
    root = _project(tmp_path, "${PREFLIGHT_URL}/items")
    monkeypatch.delenv("PREFLIGHT_URL", raising=False)
    main(["validate", "--root", str(root)])
    assert json.loads(capsys.readouterr().out)["status"] == "valid"
    with pytest.raises(ConfigurationError, match="PREFLIGHT_URL"):
        main(["validate", "--root", str(root), "--profile", "live"])
    operations = json.loads((root / "config/operations.json").read_text())
    del operations["operations"]["http.ping"]["url"]
    (root / "config/operations.json").write_text(json.dumps(operations))
    with pytest.raises(ConfigurationError, match="must define url"):
        main(["validate", "--root", str(root)])


def test_dynamic_output_and_free_business_keys_are_preserved(tmp_path, server):
    url, received = server
    root = _project(tmp_path, url)
    case = load_cases(root / "cases/demo.json")[0]
    first = replace(case.steps[0], save_as="created", request={"json": {"jsno": 3}})
    second = replace(first, id="second", order=2, save_as=None, request={
        "json": {"id": "${steps.created.body.id}", "arbitrary_business_field": True},
    })
    case = replace(case, steps=(first, second))
    with CaseRunner(root, profile="live") as runner:
        checked = runner.preflight([case])
        assert [item["step_id"] for item in checked.deferred] == ["second"]
        assert runner.run(case)["second"].output["body"]["ok"] is True
    assert received == [{"jsno": 3}, {"id": 7, "arbitrary_business_field": True}]


def test_dynamic_http_request_is_revalidated_before_next_request(tmp_path, server):
    url, received = server
    root = _project(tmp_path, url)
    case = load_cases(root / "cases/demo.json")[0]
    second = replace(case.steps[0], id="second", order=2, request={
        "json": {"id": "${steps.ping.body.id}"}, "timeout": "${steps.ping.body.bad_request}",
    })
    with CaseRunner(root, profile="live") as runner:
        with pytest.raises(ConfigurationError, match="HTTP timeout"):
            runner.run(replace(case, steps=(*case.steps, second)))
    assert len(received) == 1


def test_handler_validation_does_not_import_parent_package(tmp_path, monkeypatch, capsys):
    root = init_project(tmp_path / "project")
    package = root / "preflight_handler"
    package.mkdir()
    marker = root / "imported"
    (package / "__init__.py").write_text(f"from pathlib import Path\nPath({str(marker)!r}).touch()\n")
    (package / "actions.py").write_text("def run(**kwargs):\n    return {}\n")
    monkeypatch.syspath_prepend(str(root))
    (root / "config/operations.json").write_text(json.dumps({"operations": {
        "http.ping": {"executor": "rpc", "handler": "preflight_handler.actions:run"},
    }}))
    document = _document(root)
    document["cases"][0]["steps"][0].update(executor="rpc", expect=None)
    path = _save(root, document)
    before = _tree(root)
    main(["validate", str(path), "--root", str(root), "--profile", "live"])
    assert json.loads(capsys.readouterr().out)["data"]["deferred"]
    assert _tree(root) == before
    assert not marker.exists()


def test_custom_executor_and_injected_handler_accept_business_request(tmp_path):
    root = init_project(tmp_path / "project")
    case = load_cases(root / "cases/demo.json")[0]

    class Business(Executor):
        def execute(self, operation_name, operation, request, context):
            return ExecutionResult("http", operation_name, request)

    case = replace(case, steps=(replace(case.steps[0], request={"custom": 1}, expect={"$.custom": 1}),))
    with CaseRunner(root, profile="live", executors={"http": Business()}) as runner:
        assert runner.run(case)["ping"].output == {"custom": 1}
    (root / "config/operations.json").write_text(json.dumps({"operations": {
        "http.ping": {"executor": "rpc"},
    }}))
    with NotebookSession(root, profile="live", rpc_handlers={
        "http.ping": lambda **kwargs: kwargs["request"],
    }) as session:
        assert session.run_step(executor="rpc", operation="http.ping", request={"custom": 2}).output == {"custom": 2}


@pytest.mark.parametrize("backend", ["file", "sqlite"])
def test_missing_baseline_is_checked_without_creating_storage(tmp_path, backend):
    root = init_project(tmp_path / "project")
    (root / "config/runtime.json").write_text(json.dumps({
        "default_profile": "offline-strict", "snapshot_backend": backend,
    }))
    (root / "config/snapshots.json").write_text(json.dumps({
        "profiles": {"default": {"response_default": {}}},
    }))
    case = load_cases(root / "cases/demo.json")[0]
    case = replace(case, steps=(replace(case.steps[0], snapshot={}),))
    config = ProjectConfig(root)
    before = _tree(root)
    with pytest.raises(ConfigurationError, match="baseline does not exist"):
        preflight([case], config, config.resolve_run())
    assert _tree(root) == before


@pytest.mark.parametrize("change", [
    {"retrise": True}, {"order": 1.5}, {"order": True}, {"request": []},
    {"expect": {"checks": [{"path": "$.status_code"}]}},
    {"expect": {"checks": [{"path": "$..status_code", "equals": 200}]}},
    {"expect": {"checks": []}},
    {"mock": {"respnose": {}}}, {"snapshot": {"sleect": "$"}},
])
def test_strict_case_contract_rejects_silent_mistakes(tmp_path, change):
    root = init_project(tmp_path / "project")
    document = _document(root)
    document["cases"][0]["steps"][0].update(change)
    with pytest.raises(ContractError):
        load_cases(_save(root, document))


@pytest.mark.parametrize("sheet,column", [("cases", "enabeld"), ("steps", "requset")])
def test_excel_unknown_columns_are_rejected(tmp_path, sheet, column):
    root = init_project(tmp_path / "project")
    path = root / "cases/demo.xlsx"
    workbook = load_workbook(path)
    workbook[sheet].cell(1, workbook[sheet].max_column + 1, column)
    workbook.save(path)
    workbook.close()
    with pytest.raises(ContractError, match="unknown columns"):
        workbook_document(path)


def test_runtime_assertion_cannot_ignore_misspelled_operator():
    with pytest.raises(ContractError, match="unknown expectation"):
        assert_expectations({"status_code": 200}, {"path": "$.status_code", "equal": 418})


def test_environment_override_is_checked_only_when_inherited(tmp_path, server, monkeypatch):
    url, received = server
    monkeypatch.delenv("UNUSED_ENDPOINT", raising=False)
    root = _project(tmp_path, "${UNUSED_ENDPOINT}")
    case = load_cases(root / "cases/demo.json")[0]
    case = replace(case, steps=(replace(case.steps[0], request={"url": url}),))
    with CaseRunner(root, profile="live") as runner:
        runner.run(case)
    assert len(received) == 1


def test_database_permission_error_in_later_step_prevents_http(tmp_path, server):
    url, received = server
    root = _project(tmp_path, url)
    path = root / "config/operations.json"
    operations = json.loads(path.read_text())
    operations["operations"]["db.write"] = {
        "executor": "database", "statement": "INSERT INTO items VALUES (1)", "write": True,
    }
    path.write_text(json.dumps(operations))
    (root / "config/runtime.json").write_text(json.dumps({
        "database": {"connections": {"default": {"driver": "sqlite"}}},
        "allow_db_write": False,
    }))
    case = load_cases(root / "cases/demo.json")[0]
    later = replace(case.steps[0], id="later", order=2, executor="database", operation="db.write")
    with CaseRunner(root, profile="live") as runner:
        with pytest.raises(ConfigurationError, match="database write blocked"):
            runner.run(replace(case, steps=(*case.steps, later)))
    assert received == []


def test_source_selection_prefers_workbook_only_when_it_is_selected(tmp_path):
    root = init_project(tmp_path / "project")
    document = _document(root)
    document["cases"][0]["id"] = "json.only"
    path = _save(root, document)
    assert load_project_cases(path, write_compiled=False)[0].id == "json.only"
    assert load_project_cases(root / "cases", write_compiled=False)[0].id == "http.demo"


@pytest.mark.parametrize("reference", [
    "does_not_exist_preflight:run", "easytest.examples.handlers:no_such_function",
    "easytest.examples.handlers.run",
])
def test_invalid_later_handler_prevents_first_http_request(tmp_path, server, reference):
    url, received = server
    root = _project(tmp_path, url)
    path = root / "config/operations.json"
    operations = json.loads(path.read_text())
    operations["operations"]["rpc.later"] = {"executor": "rpc", "handler": reference}
    path.write_text(json.dumps(operations))
    case = load_cases(root / "cases/demo.json")[0]
    second = replace(case.steps[0], id="second", order=2, executor="rpc", operation="rpc.later")
    with CaseRunner(root, profile="live") as runner:
        with pytest.raises(ConfigurationError, match="handler"):
            runner.run(replace(case, steps=(*case.steps, second)))
    assert received == []


def test_sqlite_existing_baseline_validation_does_not_touch_files(tmp_path):
    root = init_project(tmp_path / "project")
    (root / "config/runtime.json").write_text(json.dumps({
        "default_profile": "offline-strict", "snapshot_backend": "sqlite",
    }))
    (root / "config/snapshots.json").write_text(json.dumps({
        "profiles": {"default": {"response_default": {}}},
    }))
    case = load_cases(root / "cases/demo.json")[0]
    case = replace(case, steps=(replace(case.steps[0], snapshot={}),))
    with CaseRunner(root, run_mode="write") as runner:
        runner.run(case)
    before = _tree(root)
    config = ProjectConfig(root)
    assert preflight([case], config, config.resolve_run(run_mode="read")).step_count == 1
    assert _tree(root) == before


def test_sqlite_preflight_defers_while_snapshot_store_is_locked(tmp_path):
    root = init_project(tmp_path / "project")
    (root / "config/runtime.json").write_text(json.dumps({
        "default_profile": "offline-strict",
        "snapshot_backend": "sqlite",
    }))
    (root / "config/snapshots.json").write_text(json.dumps({
        "profiles": {"default": {"response_default": {}}},
    }))
    case = load_cases(root / "cases/demo.json")[0]
    case = replace(case, steps=(replace(case.steps[0], snapshot={}),))
    with CaseRunner(root, run_mode="write") as runner:
        runner.run(case)
    store = SqliteSnapshotStore(root / ".easytest/snapshots.db")
    config = ProjectConfig(root)

    with store.lock:
        checked = preflight(
            [case],
            config,
            config.resolve_run(run_mode="read"),
        )

    assert checked.deferred
    assert "pending SQLite WAL" in checked.deferred[0]["reason"]


@pytest.mark.parametrize("field", ["document", "case"])
def test_unknown_container_fields_are_rejected(tmp_path, field):
    root = init_project(tmp_path / "project")
    document = _document(root)
    target = document if field == "document" else document["cases"][0]
    target["unknown_field"] = True
    with pytest.raises(ContractError, match="unknown fields"):
        load_cases(_save(root, document))


def test_cli_reuses_successful_batch_preflight_for_each_case(
    tmp_path,
    monkeypatch,
    capsys,
):
    root = init_project(tmp_path / "project")
    path = root / "cases/demo.json"
    document = json.loads(path.read_text())
    document["cases"].append({**document["cases"][0], "id": "second"})
    path.write_text(json.dumps(document))
    original = CaseRunner.preflight
    calls = []

    def tracked(self, cases, **kwargs):
        selected = list(cases)
        calls.append((len(selected), kwargs.get("reuse_for_run", False)))
        return original(self, selected, **kwargs)

    monkeypatch.setattr(CaseRunner, "preflight", tracked)
    main(["run", str(path), "--root", str(root), "--no-report"])
    assert json.loads(capsys.readouterr().out)["status"] == "passed"
    assert calls == [(2, True)]
